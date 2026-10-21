"""Excel workbook builder for FFIEC Call Report exports.

Income statement (Schedule RI) items are reported year-to-date. The workbook keeps the
as-reported YTD figures as hardcoded inputs and derives discrete quarters with live Excel
formulas (current YTD less prior-quarter YTD; Q1 equals YTD), so every derived number is
traceable to a filed figure.

House style: Segoe UI 10, no gridlines, navy #1F3864 headers, blue inputs, black formulas,
green cross-sheet links, $ in thousands, comments for annotations, no em-dashes.
"""
from __future__ import annotations

import re
from datetime import datetime
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
BLUE, BLACK, GREEN, GREY = "0000FF", "000000", "008000", "808080"
FONT = "Segoe UI"
FMT_USD = '_($* #,##0_);_($* (#,##0);_($* "-"??_);_(@_)'
FMT_NUM = '_(* #,##0_);_(* (#,##0);_(* "-"??_);_(@_)'
FMT_PCT = '0.00%'
FMT_RATIO_PCT_UNITS = '0.00'
FMT_FACTOR = '0.00'
AUTHOR = "FFIEC MCP"

NAVY_FILL = PatternFill("solid", fgColor=NAVY)
HELPER_FILL = PatternFill("solid", fgColor="D9D9D9")
TOP_BORDER = Border(top=Side(style="thin"))
FIRST_COL = 5  # column E is the first period column on statement sheets


# ---------------------------------------------------------------- statement templates
# ("item", key, line, label, [candidate MDRMs], style)
# ("formula", key, line, label, [(sign, key), ...], style)
# ("section", title) / ("blank",) / ("check", label, [(sign, key), ...])

BALANCE_SHEET = [
    ("section", "BALANCE SHEET (Schedule RC), period-end balances"),
    ("item", "bs_1a", "1.a", "Noninterest-bearing balances and currency and coin", ["RCFD0081", "RCON0081"], ""),
    ("item", "bs_1b", "1.b", "Interest-bearing balances", ["RCFD0071", "RCON0071"], ""),
    ("item", "bs_2a", "2.a", "Held-to-maturity securities", ["RCFDJJ34", "RCONJJ34", "RCFD1754", "RCON1754"], ""),
    ("item", "bs_2b", "2.b", "Available-for-sale debt securities", ["RCFD1773", "RCON1773"], ""),
    ("item", "bs_2c", "2.c", "Equity securities with readily determinable fair values", ["RCFDJA22", "RCONJA22"], ""),
    ("item", "bs_3a", "3.a", "Federal funds sold", ["RCONB987"], ""),
    ("item", "bs_3b", "3.b", "Securities purchased under agreements to resell", ["RCFDB989", "RCONB989"], ""),
    ("item", "bs_4a", "4.a", "Loans and leases held for sale", ["RCFD5369", "RCON5369"], ""),
    ("item", "bs_4b", "4.b", "Loans and leases held for investment", ["RCFDB528", "RCONB528"], ""),
    ("item", "bs_4c", "4.c", "LESS: Allowance for credit losses on loans and leases", ["RCFD3123", "RCON3123"], ""),
    ("item", "bs_4d", "4.d", "Loans and leases held for investment, net of allowance", ["RCFDB529", "RCONB529"], "subtotal"),
    ("item", "bs_5", "5", "Trading assets", ["RCFD3545", "RCON3545"], ""),
    ("item", "bs_6", "6", "Premises and fixed assets", ["RCFD2145", "RCON2145"], ""),
    ("item", "bs_7", "7", "Other real estate owned", ["RCFD2150", "RCON2150"], ""),
    ("item", "bs_8", "8", "Investments in unconsolidated subsidiaries and associated companies", ["RCFD2130", "RCON2130"], ""),
    ("item", "bs_9", "9", "Direct and indirect investments in real estate ventures", ["RCFD3656", "RCON3656"], ""),
    ("item", "bs_10", "10", "Intangible assets", ["RCFD2143", "RCON2143"], ""),
    ("item", "bs_11", "11", "Other assets", ["RCFD2160", "RCON2160"], ""),
    ("item", "bs_12", "12", "Total assets", ["RCFD2170", "RCON2170"], "total"),
    ("blank",),
    ("item", "bs_13a", "13.a", "Deposits in domestic offices", ["RCON2200"], ""),
    ("item", "bs_13a1", "13.a.(1)", "   Noninterest-bearing", ["RCON6631"], ""),
    ("item", "bs_13a2", "13.a.(2)", "   Interest-bearing", ["RCON6636"], ""),
    ("item", "bs_13b", "13.b", "Deposits in foreign offices, Edge and Agreement subsidiaries, and IBFs", ["RCFN2200"], ""),
    ("formula", "bs_dep", "13", "Total deposits", [(1, "bs_13a"), (1, "bs_13b")], "subtotal"),
    ("item", "bs_14a", "14.a", "Federal funds purchased", ["RCONB993"], ""),
    ("item", "bs_14b", "14.b", "Securities sold under agreements to repurchase", ["RCFDB995", "RCONB995"], ""),
    ("item", "bs_15", "15", "Trading liabilities", ["RCFD3548", "RCON3548"], ""),
    ("item", "bs_16", "16", "Other borrowed money", ["RCFD3190", "RCON3190"], ""),
    ("item", "bs_19", "19", "Subordinated notes and debentures", ["RCFD3200", "RCON3200"], ""),
    ("item", "bs_20", "20", "Other liabilities", ["RCFD2930", "RCON2930"], ""),
    ("item", "bs_21", "21", "Total liabilities", ["RCFD2948", "RCON2948"], "total"),
    ("blank",),
    ("item", "bs_23", "23", "Perpetual preferred stock and related surplus", ["RCFD3838", "RCON3838"], ""),
    ("item", "bs_24", "24", "Common stock", ["RCFD3230", "RCON3230"], ""),
    ("item", "bs_25", "25", "Surplus", ["RCFD3839", "RCON3839"], ""),
    ("item", "bs_26a", "26.a", "Retained earnings", ["RCFD3632", "RCON3632"], ""),
    ("item", "bs_26b", "26.b", "Accumulated other comprehensive income", ["RCFDB530", "RCONB530"], ""),
    ("item", "bs_26c", "26.c", "Other equity capital components", ["RCFDA130", "RCONA130"], ""),
    ("item", "bs_27a", "27.a", "Total bank equity capital", ["RCFD3210", "RCON3210"], "subtotal"),
    ("item", "bs_27b", "27.b", "Noncontrolling (minority) interests in consolidated subsidiaries", ["RCFD3000", "RCON3000"], ""),
    ("item", "bs_28", "28", "Total equity capital", ["RCFDG105", "RCONG105"], "total"),
    ("item", "bs_29", "29", "Total liabilities and equity capital", ["RCFD3300", "RCON3300"], "total"),
    ("check", "Check: total assets less total liabilities and equity capital (should be 0)", [(1, "bs_12"), (-1, "bs_29")]),
]

INCOME_STATEMENT = [
    ("item", "is_1a", "1.a.(3)", "Interest and fee income on loans", ["RIAD4010"], ""),
    ("item", "is_1b", "1.b", "Income from lease financing receivables", ["RIAD4065"], ""),
    ("item", "is_1c", "1.c", "Interest income on balances due from depository institutions", ["RIAD4115"], ""),
    ("item", "is_1d1", "1.d.(1)", "U.S. Treasury and U.S. Government agency securities (excl. MBS)", ["RIADB488"], ""),
    ("item", "is_1d2", "1.d.(2)", "Mortgage-backed securities", ["RIADB489"], ""),
    ("item", "is_1d3", "1.d.(3)", "All other securities", ["RIAD4060"], ""),
    ("item", "is_1e", "1.e", "Interest income from trading assets", ["RIAD4069"], ""),
    ("item", "is_1f", "1.f", "Federal funds sold and securities purchased under agreements to resell", ["RIAD4020"], ""),
    ("item", "is_1g", "1.g", "Other interest income", ["RIAD4518"], ""),
    ("item", "is_1h", "1.h", "Total interest income", ["RIAD4107"], "total"),
    ("blank",),
    ("item", "is_2a1a", "2.a.(1)(a)", "Interest on transaction accounts", ["RIAD4508"], ""),
    ("item", "is_2a1b1", "2.a.(1)(b)(1)", "Interest on savings deposits (incl. MMDAs)", ["RIAD0093"], ""),
    ("item", "is_2a1b2", "2.a.(1)(b)(2)", "Interest on time deposits of $250,000 or less", ["RIADHK03", "RIADA517"], ""),
    ("item", "is_2a1b3", "2.a.(1)(b)(3)", "Interest on time deposits of more than $250,000", ["RIADHK04", "RIADA518"], ""),
    ("item", "is_2a2", "2.a.(2)", "Interest on deposits in foreign offices", ["RIAD4172"], ""),
    ("formula", "is_2a", "2.a", "Total interest on deposits",
     [(1, "is_2a1a"), (1, "is_2a1b1"), (1, "is_2a1b2"), (1, "is_2a1b3"), (1, "is_2a2")], "subtotal"),
    ("item", "is_2b", "2.b", "Federal funds purchased and securities sold under agreements to repurchase", ["RIAD4180"], ""),
    ("item", "is_2c", "2.c", "Interest on trading liabilities and other borrowed money", ["RIAD4185"], ""),
    ("item", "is_2d", "2.d", "Interest on subordinated notes and debentures", ["RIAD4200"], ""),
    ("item", "is_2e", "2.e", "Total interest expense", ["RIAD4073"], "total"),
    ("blank",),
    ("item", "is_3", "3", "Net interest income", ["RIAD4074"], "total"),
    ("item", "is_4", "4", "Provision for credit losses", ["RIADJJ33", "RIAD4230"], ""),
    ("blank",),
    ("item", "is_5a", "5.a", "Income from fiduciary activities", ["RIAD4070"], ""),
    ("item", "is_5b", "5.b", "Service charges on deposit accounts", ["RIAD4080"], ""),
    ("item", "is_5c", "5.c", "Trading revenue", ["RIADA220"], ""),
    ("formula", "is_5oth", "5.d-5.l", "All other noninterest income (derived)",
     [(1, "is_5m"), (-1, "is_5a"), (-1, "is_5b"), (-1, "is_5c")], ""),
    ("item", "is_5m", "5.m", "Total noninterest income", ["RIAD4079"], "subtotal"),
    ("item", "is_6a", "6.a", "Realized gains (losses) on held-to-maturity securities", ["RIAD3521"], ""),
    ("item", "is_6b", "6.b", "Realized gains (losses) on available-for-sale debt securities", ["RIAD3196"], ""),
    ("blank",),
    ("item", "is_7a", "7.a", "Salaries and employee benefits", ["RIAD4135"], ""),
    ("item", "is_7b", "7.b", "Expenses of premises and fixed assets", ["RIAD4217"], ""),
    ("item", "is_7c1", "7.c.(1)", "Goodwill impairment losses", ["RIADC216"], ""),
    ("item", "is_7c2", "7.c.(2)", "Amortization expense and impairment losses for other intangibles", ["RIADC232"], ""),
    ("item", "is_7d", "7.d", "Other noninterest expense", ["RIAD4092"], ""),
    ("item", "is_7e", "7.e", "Total noninterest expense", ["RIAD4093"], "subtotal"),
    ("blank",),
    ("item", "is_8a", "8.a", "Income before change in net unrealized holding gains on equity securities", ["RIADHT69"], ""),
    ("item", "is_8b", "8.b", "Change in net unrealized holding gains (losses) on equity securities", ["RIADHT70"], ""),
    ("item", "is_8c", "8.c", "Income before applicable income taxes and discontinued operations", ["RIADHT71", "RIAD4301"], "subtotal"),
    ("item", "is_9", "9", "Applicable income taxes", ["RIAD4302"], ""),
    ("item", "is_10", "10", "Income before discontinued operations", ["RIAD4300"], ""),
    ("item", "is_11", "11", "Discontinued operations, net of applicable income taxes", ["RIADFT28"], ""),
    ("item", "is_12", "12", "Net income attributable to bank and noncontrolling interests", ["RIADG104"], ""),
    ("item", "is_13", "13", "LESS: Net income attributable to noncontrolling interests", ["RIADG103"], ""),
    ("item", "is_14", "14", "Net income attributable to bank", ["RIAD4340"], "total"),
    ("check", "Check: total interest income less total interest expense less net interest income (should be 0)",
     [(1, "is_1h"), (-1, "is_2e"), (-1, "is_3")]),
]

CAPITAL_ITEMS = [
    ("item", "cap_lev", "RC-R I.31", "Tier 1 leverage ratio (as reported)", ["RCFA7204", "RCOA7204"], "pct"),
    ("item", "cap_cet1", "RC-R I.49", "Common equity tier 1 capital ratio (as reported)", ["RCFAP793", "RCOAP793"], "pct"),
    ("item", "cap_t1", "RC-R I.50", "Tier 1 capital ratio (as reported)", ["RCFA7206", "RCOA7206"], "pct"),
    ("item", "cap_tot", "RC-R I.51", "Total capital ratio (as reported)", ["RCFA7205", "RCOA7205"], "pct"),
    ("item", "fte", "RI M.5", "Full-time equivalent employees at period end (count)", ["RIAD4150"], "count"),
]


# ---------------------------------------------------------------- helpers

def quarter_of(period: str) -> tuple[int, int]:
    m, _, y = period.split("/")
    return int(y), int(m) // 3


def qlabel(period: str) -> str:
    y, q = quarter_of(period)
    return f"Q{q} {y}"


def prior_quarter(period: str) -> str | None:
    y, q = quarter_of(period)
    if q == 1:
        return None
    return {1: f"03/31/{y}", 2: f"06/30/{y}", 3: f"09/30/{y}"}[q - 1]


def parse_value(raw: str | None) -> tuple[Any, str]:
    """Return (value, kind) with kind in num / pct / text / none."""
    if raw is None:
        return None, "none"
    s = str(raw).strip()
    if s == "":
        return None, "none"
    if s.endswith("%"):
        try:
            return float(s[:-1].replace(",", "")) / 100, "pct"
        except ValueError:
            return s, "text"
    try:
        f = float(s.replace(",", ""))
        return (int(f) if f.is_integer() else f), "num"
    except ValueError:
        return s, "text"


def sheet_title(name: str, suffix: str, used: set[str]) -> str:
    base = re.sub(r",?\s*(NATIONAL ASSOCIATION|N\.A\.|NA)$", "", name.strip(), flags=re.I)
    base = re.sub(r"[\[\]:*?/\\']", "", base).strip()
    room = 31 - len(suffix) - 1
    title = f"{base[:room].strip()} {suffix}".strip()
    n = 2
    while title.lower() in used:
        title = f"{base[:room - 3].strip()} {n} {suffix}".strip()
        n += 1
    used.add(title.lower())
    return title


def natural_key(s: str) -> list:
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s or "")]


SCHEDULE_ORDER = ["RI", "RIA", "RIBI", "RIBII", "RIC", "RID", "RIE", "RC", "RCA", "RCB", "RCCI", "RCCII",
                  "RCD", "RCEI", "RCEII", "RCF", "RCG", "RCH", "RCI", "RCK", "RCL", "RCM", "RCN", "RCO",
                  "RCP", "RCQ", "RCRI", "RCRII", "RCS", "RCT", "RCV", "CI", "ENT", "NARR"]


def schedule_key(s: str) -> tuple:
    return (SCHEDULE_ORDER.index(s) if s in SCHEDULE_ORDER else len(SCHEDULE_ORDER), s)


# ---------------------------------------------------------------- writer

class _Sheet:
    """Thin wrapper applying house style to every cell written."""

    def __init__(self, ws):
        self.ws = ws
        ws.sheet_view.showGridLines = False

    def put(self, r: int, c: int, v: Any = None, *, color: str = BLACK, bold: bool = False, fmt: str | None = None,
            italic: bool = False, size: int = 10, fill: PatternFill | None = None, align: str | None = None,
            border: Border | None = None, wrap: bool = False):
        cell = self.ws.cell(row=r, column=c)
        if v is not None:
            cell.value = v
        cell.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = fill
        if align or wrap:
            cell.alignment = Alignment(horizontal=align, vertical="top" if wrap else None, wrap_text=wrap)
        if border:
            cell.border = border
        return cell

    def header_row(self, r: int, values: list, start_col: int = 1, last_col: int | None = None):
        last_col = last_col or start_col + len(values) - 1
        for c in range(start_col, last_col + 1):
            v = values[c - start_col] if c - start_col < len(values) else None
            self.put(r, c, v, color="FFFFFF", bold=True, fill=NAVY_FILL,
                     align="center" if c >= FIRST_COL else None)

    def section(self, r: int, title: str, last_col: int):
        for c in range(1, last_col + 1):
            self.put(r, c, title if c == 1 else None, color="FFFFFF", bold=True, fill=NAVY_FILL)

    def comment(self, r: int, c: int, text: str):
        cm = Comment(text, AUTHOR)
        cm.width, cm.height = 320, 140
        self.ws.cell(row=r, column=c).comment = cm

    def widths(self, widths: dict[int, float]):
        for c, w in widths.items():
            self.ws.column_dimensions[get_column_letter(c)].width = w


def _pick(codes: list[str], data: dict[str, str]) -> str | None:
    for c in codes:
        if c in data and str(data[c]).strip() != "":
            return c
    return None


# ---------------------------------------------------------------- bank statement sheet

def build_bank_sheet(wb: Workbook, bank: dict, columns: list[str], helpers: set[str], title: str,
                     retrieved: str) -> dict:
    """Write the statements sheet for one bank. Returns row map for cross-sheet links."""
    ws = wb.create_sheet(title)
    s = _Sheet(ws)
    ncol = FIRST_COL + len(columns) - 1
    col_of = {p: FIRST_COL + i for i, p in enumerate(columns)}
    data = bank["periods"]  # period -> {mdrm: raw}

    loc = f", {bank.get('city', '')}, {bank.get('state', '')}" if bank.get("city") else ""
    s.put(1, 1, bank["name"], bold=True, size=14)
    s.put(2, 1, f"ID RSSD {bank['rssd']}  |  FDIC cert {bank.get('cert') or 'n/a'}  |  "
                f"FFIEC {bank.get('form') or 'n/a'}{loc}", color=GREY)
    s.put(3, 1, "$ in thousands unless noted. Blue = as filed; black = formula.", italic=True)
    s.put(4, 1, f"Source: FFIEC CDR Public Web Service, Call Report (SDF facsimile), retrieved {retrieved}, "
                f"https://cdr.ffiec.gov/public/", color=GREY)
    s.header_row(6, ["Line", "Item", "MDRM", "Basis"] + [qlabel(p) + ("*" if p in helpers else "") for p in columns],
                 last_col=ncol)
    for p in columns:
        c = col_of[p]
        s.put(7, c, p, color=GREY, align="center", fill=HELPER_FILL if p in helpers else None)
        if p in helpers:
            s.comment(6, c, "Helper column: fetched only so the following quarter's discrete figures can be "
                            "derived from YTD. Not part of the requested range.")
        if bank.get("missing") and p in bank["missing"]:
            s.comment(7, c, f"No Call Report available for this period: {bank['missing'][p]}")
    s.comment(6, 4, "Basis: 'Period end' = balance at the report date. 'YTD' = year-to-date flow as filed "
                    "(Jan 1 to report date). 'Quarter' = discrete three-month figure derived by formula.")

    rows: dict[str, int] = {}
    r = 9

    def write_block(template, basis: str, key_prefix: str = "", quarterly_from: dict | None = None,
                    start_row: int = 9) -> int:
        """Write template rows. quarterly_from maps key -> YTD row to derive discrete quarters."""
        nonlocal rows
        r = start_row
        need_usd = True
        # determine which rows exist (any period has a candidate code)
        present = {t[1] for t in template if t[0] == "item" and any(_pick(t[4], data.get(p, {})) for p in columns)}
        for t in template:
            kind = t[0]
            if kind == "section":
                s.section(r, t[1], ncol)
                r += 1
                need_usd = True
                continue
            if kind == "blank":
                r += 1
                need_usd = True
                continue
            if kind == "check":
                terms = [(sg, key_prefix + k) for sg, k in t[2] if key_prefix + k in rows]
                if len(terms) < 2:
                    continue
                s.put(r, 2, t[1], italic=True, color=GREY)
                for p in columns:
                    c = col_of[p]
                    L = get_column_letter(c)
                    f = "".join(("+" if sg > 0 else "-") + f"{L}{rows[k]}" for sg, k in terms).lstrip("+")
                    s.put(r, c, f"=IFERROR({f},\"\")", fmt=FMT_NUM, italic=True, color=GREY)
                r += 2
                need_usd = True
                continue
            key, line, label, spec, style = t[1], t[2], t[3], t[4], t[5]
            if kind == "item" and key not in present:
                continue
            if kind == "formula":
                terms = [(sg, key_prefix + k) for sg, k in spec if (k in present or key_prefix + k in rows
                                                                     or any(x[0] == "formula" and x[1] == k for x in template))]
                if not any(k.replace(key_prefix, "", 1) in present for _, k in terms):
                    continue
            bold = style in ("total", "subtotal")
            border = TOP_BORDER if style == "total" else None
            s.put(r, 1, line, color=GREY)
            s.put(r, 2, label, bold=bold)
            s.put(r, 4, basis, color=GREY)
            rows[key_prefix + key] = r
            is_pct, is_count = style == "pct", style == "count"
            codes_used = set()
            first_ok = next((q for q in columns if quarterly_from is None or quarter_of(q)[1] == 1
                             or prior_quarter(q) in col_of), columns[0])
            for i, p in enumerate(columns):
                c = col_of[p]
                L = get_column_letter(c)
                leftmost = p == first_ok
                fmt = FMT_PCT if is_pct else FMT_NUM if is_count else (
                    FMT_USD if leftmost and (need_usd or bold) else FMT_NUM)
                fill = HELPER_FILL if p in helpers else None
                if kind == "formula":
                    if quarterly_from is not None and quarter_of(p)[1] != 1 and prior_quarter(p) not in col_of:
                        s.put(r, c, "n/a", color=GREY, align="right", fill=fill, border=border)
                        continue
                    refs = [(sg, k) for sg, k in terms]
                    f = "".join(("+" if sg > 0 else "-") + f"N({L}{{{k}}})" for sg, k in refs).lstrip("+")
                    s.put(r, c, ("=" + f), fmt=fmt, bold=bold, border=border, fill=fill)
                    continue
                if quarterly_from is not None:
                    ytd_row = quarterly_from[key]
                    y, q = quarter_of(p)
                    pp = prior_quarter(p)
                    if q == 1:
                        f = f"=IF(ISNUMBER({L}{ytd_row}),{L}{ytd_row},\"\")"
                    elif pp in col_of:
                        PL = get_column_letter(col_of[pp])
                        f = (f"=IF(AND(ISNUMBER({L}{ytd_row}),ISNUMBER({PL}{ytd_row})),"
                             f"{L}{ytd_row}-{PL}{ytd_row},\"\")")
                    else:
                        f = None
                    if f:
                        s.put(r, c, f, fmt=fmt, bold=bold, border=border, fill=fill)
                    else:
                        s.put(r, c, "n/a", color=GREY, align="right", fill=fill, border=border)
                    continue
                code = _pick(spec, data.get(p, {}))
                if code:
                    codes_used.add(code)
                    v, vk = parse_value(data[p][code])
                    s.put(r, c, v, color=BLUE, fmt=fmt, bold=bold, border=border, fill=fill)
                else:
                    s.put(r, c, None, border=border, fill=fill)
            if kind == "item":
                s.put(r, 3, ", ".join(sorted(codes_used)) if codes_used else spec[0], color=GREY)
                if quarterly_from is not None:
                    s.put(r, 3, ws.cell(row=quarterly_from[key], column=3).value, color=GREY)
            need_usd = False
            r += 1
        # resolve formula placeholders now that all rows in this block are known
        for rr in range(start_row, r):
            for c in range(FIRST_COL, ncol + 1):
                v = ws.cell(row=rr, column=c).value
                if isinstance(v, str) and "{" in v:
                    def sub(m):
                        k = m.group(1)
                        return str(rows[k]) if k in rows else "#DROP#"
                    nv = re.sub(r"\{(\w+)\}", sub, v)
                    nv = re.sub(r"[+-]?N\([A-Z]+#DROP#\)", "", nv).replace("=+", "=")
                    ws.cell(row=rr, column=c).value = nv if nv != "=" else None
        return r

    r = write_block(BALANCE_SHEET, "Period end", start_row=r)
    r += 1
    ytd_template = [("section", "INCOME STATEMENT (Schedule RI), YEAR-TO-DATE AS FILED")] + INCOME_STATEMENT
    ytd_start = r
    r = write_block(ytd_template, "YTD", start_row=r)
    s.comment(ytd_start, 1, "Call Report income statement items are cumulative from January 1 to the report "
                            "date. Q2 YTD = six months, Q3 YTD = nine months, Q4 YTD = full year. Do not add "
                            "quarterly YTD figures together.")
    r += 1
    ytd_rows = {k: v for k, v in rows.items() if k.startswith("is_")}
    q_template = [("section", "INCOME STATEMENT (Schedule RI), DISCRETE QUARTER (derived from YTD)")] + [
        t for t in INCOME_STATEMENT if t[0] != "check"] + [INCOME_STATEMENT[-1]]
    q_start = r
    r = write_block(q_template, "Quarter", key_prefix="q_", quarterly_from=ytd_rows, start_row=r)
    s.comment(q_start, 1, "Discrete quarter = current-quarter YTD less prior-quarter YTD of the same year; "
                          "Q1 equals Q1 YTD. Shown as 'n/a' where the prior quarter was not retrieved. "
                          "Amended filings: YTD figures reflect the latest amendment on file at retrieval.")

    # key ratios
    r += 1
    s.section(r, "KEY RATIOS AND REGULATORY CAPITAL", ncol)
    r += 1
    s.put(r, 2, "Annualization factor (12 / months in YTD period)")
    s.put(r, 4, "Input", color=GREY)
    rows["factor"] = r
    for p in columns:
        _, q = quarter_of(p)
        s.put(r, col_of[p], 4 / q, color=BLUE, fmt=FMT_FACTOR, fill=HELPER_FILL if p in helpers else None)
    r += 1

    def ratio(key, label, basis, expr, note=None):
        nonlocal r
        needed = re.findall(r"\{(\w+)\}", expr)
        if not all(k in rows for k in needed):
            return
        s.put(r, 2, label)
        s.put(r, 4, basis, color=GREY)
        rows[key] = r
        for p in columns:
            L = get_column_letter(col_of[p])
            f = re.sub(r"\{(\w+)\}", lambda m: f"{L}{rows[m.group(1)]}", expr)
            s.put(r, col_of[p], f"=IFERROR({f},\"\")", fmt=FMT_PCT, fill=HELPER_FILL if p in helpers else None)
        if note:
            s.comment(r, 2, note)
        r += 1

    ratio("roa_ytd", "Return on assets, annualized (YTD net income / period-end assets)", "YTD",
          "{is_14}*{factor}/{bs_12}", "Uses period-end total assets, not average assets, so it will differ "
          "from the UBPR ROA (which uses average assets).")
    ratio("roa_q", "Return on assets, annualized (quarter net income x 4 / period-end assets)", "Quarter",
          "{q_is_14}*4/{bs_12}")
    ratio("roe_ytd", "Return on bank equity, annualized (YTD)", "YTD", "{is_14}*{factor}/{bs_27a}")
    ratio("nii_ta", "Net interest income / total assets, annualized (YTD)", "YTD", "{is_3}*{factor}/{bs_12}")
    ratio("eff_ytd", "Efficiency ratio (noninterest expense / (NII + noninterest income)), YTD", "YTD",
          "{is_7e}/({is_3}+{is_5m})")
    ratio("eff_q", "Efficiency ratio, discrete quarter", "Quarter", "{q_is_7e}/({q_is_3}+{q_is_5m})")
    ratio("ltd", "Loans held for investment / total deposits", "Period end", "{bs_4b}/{bs_dep}")
    ratio("acl", "Allowance / loans held for investment", "Period end", "{bs_4c}/{bs_4b}")
    ratio("eq_ta", "Total bank equity capital / total assets", "Period end", "{bs_27a}/{bs_12}")
    r = write_block(CAPITAL_ITEMS, "Period end", start_row=r)

    s.widths({1: 12, 2: 62, 3: 18, 4: 11, **{c: 15 for c in range(FIRST_COL, ncol + 1)}})
    return {"title": title, "rows": rows, "col_of": col_of}


# ---------------------------------------------------------------- detail sheet

def build_detail_sheet(wb: Workbook, bank: dict, columns: list[str], helpers: set[str], title: str,
                       select: Callable[[str, str], bool], label_of: Callable[[str], str | None]) -> int:
    ws = wb.create_sheet(title)
    s = _Sheet(ws)
    data, desc, basis_map = bank["periods"], bank["desc"], bank.get("basis", {})
    codes = sorted({c for p in columns for c in data.get(p, {}) if select(c, desc.get(c, ("", "", ""))[1])},
                   key=lambda c: (schedule_key(desc.get(c, ("", "", ""))[1]), natural_key(desc.get(c, ("", "", ""))[2]), c))
    n = len(columns)
    a0 = 7                      # first as-reported column
    q0 = a0 + n + 1             # first discrete-quarter column
    last = q0 + n - 1
    s.put(1, 1, f"{bank['name']}: Call Report detail", bold=True, size=14)
    s.put(2, 1, "$ in thousands unless noted. Blue = as filed; black = formula. Basis flag taken from the "
                "filing's XBRL context (duration = YTD flow, instant = period end).", italic=True)
    s.header_row(4, ["Schedule", "Line", "MDRM", "Description (as filed)", "MDRM dictionary name", "Basis"], 1, 6)
    s.put(4, a0, "AS FILED", color="FFFFFF", bold=True, fill=NAVY_FILL)
    for c in range(a0 + 1, a0 + n):
        s.put(4, c, None, fill=NAVY_FILL)
    s.put(4, q0, "DISCRETE QUARTER (YTD flows only, derived)", color="FFFFFF", bold=True, fill=NAVY_FILL)
    for c in range(q0 + 1, last + 1):
        s.put(4, c, None, fill=NAVY_FILL)
    for i, p in enumerate(columns):
        for c0 in (a0, q0):
            s.put(5, c0 + i, qlabel(p) + ("*" if p in helpers else ""), bold=True, align="center",
                  fill=HELPER_FILL if p in helpers else None)
    r = 6
    prev_sched = None
    for code in codes:
        d, sched, line = desc.get(code, ("", "", ""))
        if sched != prev_sched:
            s.section(r, f"Schedule {sched}", last)
            r += 1
            prev_sched = sched
        pt = basis_map.get(code) or ("duration" if code.startswith("RIAD") else "instant")
        vals = {p: parse_value(data.get(p, {}).get(code)) for p in columns}
        is_text = any(k == "text" for v, k in vals.values())
        basis = "Text" if is_text else ("YTD" if pt == "duration" else "Period end")
        s.put(r, 1, sched, color=GREY)
        s.put(r, 2, line, color=GREY)
        s.put(r, 3, code, color=GREY)
        s.put(r, 4, d)
        s.put(r, 5, label_of(code) or "", color=GREY)
        s.put(r, 6, basis, color=GREY)
        for i, p in enumerate(columns):
            v, k = vals[p]
            fmt = FMT_PCT if k == "pct" else FMT_NUM
            s.put(r, a0 + i, v, color=BLUE, fmt=fmt, fill=HELPER_FILL if p in helpers else None)
            if basis == "YTD":
                L = get_column_letter(a0 + i)
                _, q = quarter_of(p)
                pp = prior_quarter(p)
                if q == 1:
                    f = f"=IF(ISNUMBER({L}{r}),{L}{r},\"\")"
                elif pp in columns:
                    PL = get_column_letter(a0 + columns.index(pp))
                    f = f"=IF(AND(ISNUMBER({L}{r}),ISNUMBER({PL}{r})),{L}{r}-{PL}{r},\"\")"
                else:
                    f = None
                s.put(r, q0 + i, f if f else "n/a", color=BLACK if f else GREY, fmt=fmt,
                      fill=HELPER_FILL if p in helpers else None, align=None if f else "right")
        r += 1
    s.widths({1: 9, 2: 9, 3: 11, 4: 60, 5: 40, 6: 10, **{c: 14 for c in range(a0, last + 1)}, a0 + n: 3})
    return len(codes)


# ---------------------------------------------------------------- workbook

def build_workbook(path: str, banks: list[dict], columns: list[str], helpers: set[str],
                   detail_select: Callable[[str, str], bool] | None, label_of: Callable[[str], str | None],
                   ubpr: dict | None = None, ubpr_fraction: dict | None = None) -> dict:
    wb = Workbook()
    notes = _Sheet(wb.active)
    notes.ws.title = "Notes"
    used = {"notes"}
    retrieved = datetime.now().strftime("%m/%d/%Y %I:%M %p")
    requested = [p for p in columns if p not in helpers]
    built = []
    for b in banks:
        info = build_bank_sheet(wb, b, columns, helpers, sheet_title(b["name"], "", used), retrieved)
        info["bank"] = b
        if detail_select:
            info["detail_rows"] = build_detail_sheet(wb, b, columns, helpers, sheet_title(b["name"], "Detail", used),
                                                     detail_select, label_of)
        built.append(info)

    # comparison sheet (latest requested period)
    if len(built) > 1:
        cmp_ = _Sheet(wb.create_sheet("Compare", 1))
        latest = requested[-1]
        cmp_.put(1, 1, f"Peer comparison, {qlabel(latest)} ({latest})", bold=True, size=14)
        cmp_.put(2, 1, "$ in thousands. Green = linked from each bank's sheet.", italic=True)
        cmp_.header_row(4, ["Metric", "Basis"] + [i["bank"]["name"] for i in built], 1, 2 + len(built))
        metrics = [("bs_12", "Total assets", "Period end", "usd"), ("bs_4b", "Loans held for investment", "Period end", "num"),
                   ("bs_dep", "Total deposits", "Period end", "num"), ("bs_27a", "Total bank equity capital", "Period end", "num"),
                   ("is_3", "Net interest income", "YTD", "num"), ("q_is_3", "Net interest income", "Quarter", "num"),
                   ("is_14", "Net income attributable to bank", "YTD", "num"),
                   ("q_is_14", "Net income attributable to bank", "Quarter", "num"),
                   ("roa_ytd", "Return on assets, annualized", "YTD", "pct"), ("roa_q", "Return on assets, annualized", "Quarter", "pct"),
                   ("eff_ytd", "Efficiency ratio", "YTD", "pct"), ("ltd", "Loans / deposits", "Period end", "pct"),
                   ("cap_cet1", "CET1 ratio", "Period end", "pct"), ("cap_lev", "Tier 1 leverage ratio", "Period end", "pct")]
        r = 5
        for key, label, basis, kind in metrics:
            cmp_.put(r, 1, label)
            cmp_.put(r, 2, basis, color=GREY)
            for j, info in enumerate(built):
                c = 3 + j
                if key in info["rows"] and latest in info["col_of"]:
                    L = get_column_letter(info["col_of"][latest])
                    fmt = FMT_PCT if kind == "pct" else (FMT_USD if kind == "usd" and j == 0 else FMT_NUM)
                    cmp_.put(r, c, f"='{info['title']}'!{L}{info['rows'][key]}", color=GREEN, fmt=fmt)
            r += 1
        cmp_.widths({1: 40, 2: 11, **{3 + j: 22 for j in range(len(built))}})
        for j in range(len(built)):
            cmp_.put(4, 3 + j, None, color="FFFFFF", bold=True, fill=NAVY_FILL, wrap=True)

    # long-format data sheet for pivots
    data_ws = _Sheet(wb.create_sheet("Data"))
    hdr = ["Bank", "ID RSSD", "Period end", "Quarter", "Helper", "Statement", "Line", "MDRM", "Item", "Basis",
           "As filed ($000 or ratio)", "Discrete quarter ($000)"]
    data_ws.header_row(1, hdr, 1, len(hdr))
    for c in range(1, len(hdr) + 1):
        data_ws.ws.cell(row=1, column=c).alignment = Alignment(horizontal="left")
    r = 2
    tmpl = ([("Balance sheet", t) for t in BALANCE_SHEET if t[0] == "item"] +
            [("Income statement", t) for t in INCOME_STATEMENT if t[0] == "item"] +
            [("Regulatory capital", t) for t in CAPITAL_ITEMS])
    for b in banks:
        for stmt, t in tmpl:
            _, key, line, label, codes, style = t
            ytd_flow = stmt == "Income statement"
            vals = {}
            for p in columns:
                code = _pick(codes, b["periods"].get(p, {}))
                vals[p] = (code, *parse_value(b["periods"][p][code])) if code else (None, None, "none")
            if all(v[0] is None for v in vals.values()):
                continue
            for p in columns:
                code, v, k = vals[p]
                if code is None:
                    continue
                disc = None
                if ytd_flow and k == "num":
                    _, q = quarter_of(p)
                    pp = prior_quarter(p)
                    if q == 1:
                        disc = v
                    elif pp in vals and vals[pp][2] == "num":
                        disc = v - vals[pp][1]
                row = [b["name"], int(b["rssd"]) if str(b["rssd"]).isdigit() else b["rssd"], p, qlabel(p),
                       "Yes" if p in helpers else "", stmt, line, code, label.strip(),
                       "YTD" if ytd_flow else "Period end", v, disc]
                for c, val in enumerate(row, 1):
                    fmt = (FMT_PCT if k == "pct" else FMT_NUM) if c in (11, 12) else None
                    data_ws.put(r, c, val, color=BLUE if c in (11, 12) else BLACK, fmt=fmt)
                r += 1
    data_ws.widths({1: 34, 2: 10, 3: 12, 4: 9, 5: 7, 6: 18, 7: 11, 8: 11, 9: 60, 10: 10, 11: 20, 12: 20})
    data_ws.ws.auto_filter.ref = f"A1:L{max(r - 1, 1)}"

    # UBPR
    if ubpr:
        u = _Sheet(wb.create_sheet("UBPR"))
        ucols = requested
        u.put(1, 1, "Uniform Bank Performance Report (UBPR) concepts, as published", bold=True, size=14)
        u.put(2, 1, "Dollar items in $ thousands (converted from whole dollars). Ratios in percent units as "
                    "published (1.05 = 1.05%), except rows marked 'fraction', which UBPR stores as decimals and "
                    "are shown here as percentages. Not de-cumulated.", italic=True)
        u.header_row(4, ["Bank", "Concept", "Name (MDRM dictionary)", "Unit"] + [qlabel(p) for p in ucols], 1, 4 + len(ucols))
        u.comment(4, 3, "UBPR income-based ratios are computed on a year-to-date basis and annualized; UBPR "
                        "dollar income items are year-to-date. Balance-based ratios use averages defined in "
                        "the UBPR User's Guide. Figures here are left exactly as published by the FFIEC.")
        r = 5
        for b in banks:
            per = ubpr.get(str(b["rssd"]), {})
            concepts = sorted({c for p in ucols for c in per.get(p, {})})
            for code in concepts:
                unit = next((per[p][code][1] for p in ucols if code in per.get(p, {})), "")
                u.put(r, 1, b["name"])
                u.put(r, 2, code, color=GREY)
                u.put(r, 3, label_of(code) or "")
                frac = (ubpr_fraction or {}).get(code)
                u.put(r, 4, "$000" if unit == "USD" else "fraction" if frac else "%" if unit in ("PURE", "PERCENT")
                      else (unit or ""), color=GREY)
                if frac:
                    u.comment(r, 4, frac)
                for i, p in enumerate(ucols):
                    if code in per.get(p, {}):
                        v, unit = per[p][code]
                        v, k = parse_value(v)
                        if k == "num" and unit == "USD":
                            v = v / 1000
                        fmt = FMT_NUM if unit == "USD" else FMT_PCT if frac else FMT_RATIO_PCT_UNITS
                        u.put(r, 5 + i, v, color=BLUE, fmt=fmt)
                r += 1
        u.widths({1: 34, 2: 11, 3: 60, 4: 7, **{5 + i: 14 for i in range(len(ucols))}})

    # notes sheet
    n = notes
    n.put(1, 1, "FFIEC Call Report Export", bold=True, size=14)
    n.put(2, 1, f"Generated {retrieved}", color=GREY)
    n.header_row(4, ["Bank", "ID RSSD", "FDIC cert", "Form", "Location", "Sheet"], 1, 6)
    for c in range(1, 7):
        n.ws.cell(row=4, column=c).alignment = Alignment(horizontal="left")
    r = 5
    for info in built:
        b = info["bank"]
        for c, v in enumerate([b["name"], b["rssd"], b.get("cert"), b.get("form"),
                               f"{b.get('city', '')}, {b.get('state', '')}".strip(", "), info["title"]], 1):
            n.put(r, c, v)
        r += 1
    r += 1
    lines = [
        ("PERIODS", None),
        (f"Requested: {qlabel(requested[0])} to {qlabel(requested[-1])} ({len(requested)} quarters).", None),
        ("Columns marked * (grey) are helper quarters fetched only to derive the next quarter's discrete figures.", None),
        ("", None),
        ("YEAR-TO-DATE CONVENTION", None),
        ("Call Report income statement items (Schedule RI and other RIAD items) are reported year-to-date: "
         "Q1 = 3 months, Q2 = 6 months, Q3 = 9 months, Q4 = 12 months. Never sum YTD columns.", None),
        ("Each bank sheet shows the income statement twice: YEAR-TO-DATE AS FILED (hardcoded, blue) and "
         "DISCRETE QUARTER (formulas, black): current YTD less prior-quarter YTD of the same year; Q1 equals Q1 YTD.", None),
        ("Balance sheet items and regulatory capital ratios are period-end values and are not adjusted.", None),
        ("Annualized ratios scale YTD income by 12 / months elapsed (factor row on each bank sheet).", None),
        ("", None),
        ("CAVEATS", None),
        ("Values reflect the latest amended filing on file at retrieval; an amendment to an earlier quarter "
         "changes the next quarter's derived discrete figure.", None),
        ("Mergers during a year: the acquirer's YTD income includes the target only from the acquisition date, "
         "so derived quarters are not comparable to pre-merger periods.", None),
        ("FFIEC 051 filers report fewer items; rows with no data for a bank are omitted from its sheet.", None),
        ("MDRM codes can change with form revisions; each row lists the code(s) actually used.", None),
        ("", None),
        ("SOURCE", None),
        (f"Source: FFIEC Central Data Repository Public Web Service (SIS611), Call Report SDF facsimiles, "
         f"retrieved {retrieved}, https://cdr.ffiec.gov/public/", None),
        ("Item names: Federal Reserve Micro Data Reference Manual (MDRM), https://www.federalreserve.gov/apps/mdrm/", None),
        ("", None),
        ("LEGEND", None),
        ("Blue = as filed (hardcoded input); black = formula; green = cross-sheet link; navy = section header.", None),
    ]
    for text, _ in lines:
        if text.isupper() and text:
            n.section(r, text, 6)
        else:
            n.put(r, 1, text)
        r += 1
    n.put(r, 1, "Blue sample", color=BLUE)
    n.put(r, 2, "Black sample")
    n.put(r, 3, "Green sample", color=GREEN)
    n.widths({1: 44, 2: 12, 3: 12, 4: 8, 5: 24, 6: 34})

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.save(path)
    return {"sheets": wb.sheetnames, "detail_rows": {i["title"]: i.get("detail_rows") for i in built}}
