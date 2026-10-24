"""Offline tests: every tool runs against a mocked FFIEC API. No credentials or network needed.

Run: python tests/test_offline.py
"""
import asyncio
import base64
import json
import os
import sys
import re
import tempfile
import time
import zipfile
from datetime import datetime
from pathlib import Path

TMP = Path(tempfile.mkdtemp(prefix="ffiec_test_"))
os.environ.update(FFIEC_USER_ID="tester", FFIEC_TOKEN="Bearer test-token", FFIEC_DOWNLOAD_DIR=str(TMP))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import httpx  # noqa: E402
import server  # noqa: E402
from store import EASTERN  # noqa: E402

server.MIN_INTERVAL = 0
HEADER = '"Call Date";"Bank RSSD Identifier";"MDRM #";"Value";"Last Update";"Short Definition";"Call Schedule";"Line Number"\n'


def sdf(period: str, ni_ytd: int, nii_ytd: int) -> str:
    d = period[6:] + period[:2] + period[3:5]
    rows = [("RCON2170", 1000000, "Total assets", "RCRII", "11"), ("RCON3300", 1000000, "Total L&E", "RC", "29"),
            ("RCONB528", 600000, "Loans HFI", "RC", "4b"), ("RCON2200", 800000, "Deposits", "RC", "13a"),
            ("RCON3210", 100000, "Total bank equity", "RC", "27a"), ("RIAD4107", nii_ytd + 100, "Total interest income", "RI", "1h"),
            ("RIAD4073", 100, "Total interest expense", "RI", "2e"), ("RIAD4074", nii_ytd, "Net interest income", "RI", "3"),
            ("RIAD4340", ni_ytd, "Net income", "RIA", "4"), ("RIAD4150", 250, "FTE employees", "RI", "M5"),
            ("RCOA7204", "10.5%", "Leverage ratio", "RCRI", "31")]
    return HEADER + "".join(f'{d};480228;"{c}";"{v}";"{d}";"{t}";"{s}";"{l}"\n' for c, v, t, s, l in rows)


XBRL = b'''<xbrl xmlns="http://www.xbrl.org/2003/instance" xmlns:uc="u">
<context id="CI_1_2026-06-30"><entity/><period><instant>2026-06-30</instant></period></context>
<context id="CD_1_2026-06-30"><entity/><period><startDate>2026-01-01</startDate><endDate>2026-06-30</endDate></period></context>
<uc:UBPRE013 contextRef="CI_1_2026-06-30" unitRef="PURE" decimals="2">1.05</uc:UBPRE013>
<uc:RIAD4340 contextRef="CD_1_2026-06-30" unitRef="USD" decimals="0">3000000</uc:RIAD4340>
<uc:RIAD4150 contextRef="CI_1_2026-06-30" unitRef="NON-MONETARY" decimals="0">250</uc:RIAD4150>
<uc:UBPR7204 contextRef="CI_1_2026-06-30" unitRef="PURE" decimals="2">0.11</uc:UBPR7204></xbrl>'''
YTD = {"03/31/2026": (1000, 5000), "06/30/2026": (3000, 10500)}
SUBMISSIONS = []  # extra RetrieveFilersSubmissionDateTime rows, to simulate amendments
calls = []


def fake_get(url, headers, timeout, **kw):
    m = url.rsplit("/", 1)[1]
    calls.append((m, headers))
    req = httpx.Request("GET", url)
    J = lambda o: httpx.Response(200, json=o, request=req)  # noqa: E731
    if m in ("RetrieveReportingPeriods", "RetrieveUBPRReportingPeriods"):
        return J(["12/31/2025", "3/31/2026", "6/30/2026"])  # FFIEC does not zero-pad
    if m == "RetrievePanelOfReporters":
        return J([{"ID_RSSD": 480228, "FDICCertNumber": 3510, "Name": "TEST BANK, NATIONAL ASSOCIATION    ",
                   "State": "NC", "City": "CHARLOTTE  ", "FilingType": "041", "HasFiledForReportingPeriod": True}])
    if m == "RetrieveFilersSubmissionDateTime":
        return J(SUBMISSIONS + [{"ID_RSSD": 1, "DateTime": "7/1/2026 6:17:33 PM"}, {"ID_RSSD": 2, "DateTime": "7/9/2026 11:42:45 AM"}])
    if m == "RetrieveFacsimile":
        p = headers["reportingPeriodEndDate"]
        if headers["fiId"] == "999" or p not in YTD:
            return httpx.Response(204, request=req)
        fmt = headers["facsimileFormat"]
        body = sdf(p, *YTD[p]).encode() if fmt == "SDF" else XBRL if fmt == "XBRL" else b"%PDF-1.4"
        return J(base64.b64encode(body).decode()) if fmt != "PDF" else J(list(body))
    if m == "RetrieveUBPRXBRLFacsimile":
        return J(base64.b64encode(XBRL).decode())
    if m == "RetrieveFilersSinceDate":
        return httpx.Response(500, json={"Message": "Error Code 5002: Security Token is invalid."}, request=req)
    raise AssertionError(m)


server.httpx.get = fake_get
failures = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  {detail}"))
    if not cond:
        failures.append(name)


async def call(name, args):
    r = await server.mcp.call_tool(name, args)
    return r[1].get("result", r[1]) if isinstance(r, tuple) else json.loads(r[0].text)


async def main():
    r = await call("ffiec_reporting_periods", {})
    check("periods sorted, zero-padded", r["latest"] == "06/30/2026" and r["periods"][-1] == "12/31/2025", r)
    r = await call("ffiec_find_bank", {"query": "test bank"})
    check("find bank strips padding", r["results"][0]["City"] == "CHARLOTTE", r)
    r = await call("ffiec_recent_filers", {"period": "2026Q2", "since": "2026-07-01"})
    check("recent filers newest first", r["filers"][0]["ID_RSSD"] == 2, r)
    r = await call("ffiec_recent_filers", {"period": "2026Q2", "since": "07/01/2026", "include_times": False})
    check("token error gets a hint", "Regenerate" in r.get("error", ""), r)
    r = await call("ffiec_call_report", {"fi_id": "480228", "period": "2026Q2", "items": ["2170"]})
    check("call report item filter", r["items"][0]["mdrm"] == "RCON2170", r)
    r = await call("ffiec_call_report", {"fi_id": "480228", "period": "2026Q2", "schedule": "RC-R part II"})
    check("schedule filter ignores hyphens", [i["mdrm"] for i in r["items"]] == ["RCON2170"], r)
    r = await call("ffiec_bank_summary", {"fi_ids": ["480228", "999"], "period": "2026Q2"})
    b = r["banks"]["480228"]
    check("summary parses % ratios", b["tier1_leverage_ratio_pct"] == 10.5, b)
    check("summary reports missing bank", "error" in r["banks"]["999"], r)
    r = await call("ffiec_ubpr", {"fi_id": "480228", "period": "2026Q2", "search": "average assets"})
    check("ubpr names attached and searchable", r["facts"] and r["facts"][0]["concept"] == "UBPRE013", r)
    r = await call("ffiec_lookup_items", {"query": "ROA", "prefix": "UBPR"})
    check("lookup synonyms", any(x["code"] == "UBPRE013" for x in r["results"]), r)
    r = await call("ffiec_lookup_items", {"query": "2170", "limit": 3})
    check("lookup bare item number", r["results"][0]["code"] == "UBPR2170", r)

    out = TMP / "export.xlsx"
    r = await call("ffiec_export_excel", {"fi_ids": ["480228"], "periods": ["2026Q2"], "detail_items": ["RIAD4340", "RIAD4150"],
                                          "ubpr_concepts": ["UBPRE013"], "output_path": str(out)})
    check("export ran", r.get("file") == str(out) and r["helper_quarters"] == ["Q1 2026"], r)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb[r["sheets"][1]]
    found = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        label, basis = row[1].value, row[3].value
        if label == "Net income attributable to bank":
            found[basis] = (row[4].value, row[5].value)
    check("YTD net income as filed", found.get("YTD") == (1000, 3000), found)
    q = found.get("Quarter", ("", ""))
    check("discrete Q1 = YTD", str(q[0]).startswith("=IF(ISNUMBER(E"), q)
    check("discrete Q2 = YTD less Q1 YTD", "F" in str(q[1]) and "-E" in str(q[1]), q)
    det = wb[[s for s in wb.sheetnames if s.endswith("Detail")][0]]
    bases = {row[2].value: row[5].value for row in det.iter_rows(min_row=6) if row[2].value}
    check("XBRL basis flags (employees are period end)", bases.get("RIAD4150") == "Period end" and bases.get("RIAD4340") == "YTD", bases)
    check("notes, data, UBPR sheets", {"Notes", "Data", "UBPR"} <= set(wb.sheetnames), wb.sheetnames)

    # resumable export: a zero time budget forces one download per call
    server.TIME_BUDGET = 0
    YTD["12/31/2025"] = (4000, 20000)
    YTD["09/30/2025"] = (3000, 15000)
    args = {"fi_ids": ["480228"], "periods": ["2025Q4"], "output_path": str(TMP / "resume.xlsx")}
    statuses = []
    for _ in range(5):
        r = await call("ffiec_export_excel", args)
        statuses.append(r.get("status"))
        if r.get("status") == "complete":
            break
    check("export resumes across calls", statuses == ["in_progress", "complete"], statuses)
    server.TIME_BUDGET = 40
    r = await call("ffiec_export_excel", {"fi_ids": ["999"], "periods": ["2026Q2"], "output_path": str(TMP / "none.xlsx")})
    check("missing filings do not loop", "error" in r and "No data" in r["error"], r)
    n_before = len(calls)
    await call("ffiec_export_excel", {"fi_ids": ["999"], "periods": ["2026Q2"], "output_path": str(TMP / "none.xlsx")})
    check("missing filings remembered (no repeat API calls)", len(calls) == n_before, len(calls) - n_before)

    r = await call("ffiec_ubpr", {"fi_id": "480228", "period": "2026Q2", "concepts": ["UBPR7204"]})
    check("UBPR fraction ratio flagged", r["facts"] and "fraction" in r["facts"][0].get("unit_note", ""), r)

    docs = server._documents_dir()
    check("dir: unexpanded ${DOCUMENTS} default", server._resolve_dir("${DOCUMENTS}/FFIEC") == docs / "FFIEC")
    check("dir: empty -> Documents/FFIEC", server._resolve_dir("") == docs / "FFIEC")
    check("dir: relative -> Documents/FFIEC", server._resolve_dir("FFIEC") == docs / "FFIEC")
    check("dir: unknown variable -> default", server._resolve_dir("${FOO}/x") == docs / "FFIEC")
    check("dir: absolute kept", server._resolve_dir(str(TMP)) == TMP)

    # ---- amendment detection
    db = server._db()
    period = "06/30/2026"
    r = await call("ffiec_call_report", {"fi_id": "480228", "period": period, "items": ["RIAD4340"]})
    check("cached before amendment", r["items"][0]["value"] == "3000", r)
    YTD[period] = (3500, 10500)  # the bank amends its net income
    stamp = datetime.fromtimestamp(time.time() + 3600, EASTERN).strftime("%-m/%-d/%Y %-I:%M:%S %p")
    SUBMISSIONS.append({"ID_RSSD": 480228, "DateTime": stamp})
    n = len(calls)
    r = await call("ffiec_call_report", {"fi_id": "480228", "period": period, "items": ["RIAD4340"]})
    check("amendment not seen before the check is due", r["items"][0]["value"] == "3000" and len(calls) == n, r)
    db._w("UPDATE checks SET checked_at=? WHERE period=?", (time.time() - server.CHECK_TTL - 1, period))
    n = len(calls)
    r = await call("ffiec_call_report", {"fi_id": "480228", "period": period, "items": ["RIAD4340"]})
    methods = [c[0] for c in calls[n:]]
    check("amended filing re-downloaded", r["items"][0]["value"] == "3500", r)
    check("amendment check costs one call", methods == ["RetrieveFilersSubmissionDateTime", "RetrieveFacsimile"], methods)
    check("check sent period and a since-date", calls[n][1]["reportingPeriodEndDate"] == period
          and re.fullmatch(r"\d\d/\d\d/\d{4}", calls[n][1]["lastUpdateDateTime"]), calls[n][1])
    n = len(calls)
    await call("ffiec_bank_summary", {"fi_ids": ["480228"], "period": period})
    check("no calls once re-downloaded", len(calls) == n, [c[0] for c in calls[n:]])
    SUBMISSIONS.clear()

    # a failed check never blocks cached data, and backs off
    db._w("UPDATE checks SET checked_at=? WHERE period=?", (time.time() - server.CHECK_TTL - 1, period))
    real_get = server.httpx.get
    server.httpx.get = lambda url, **kw: (calls.append((url.rsplit("/", 1)[1], kw["headers"])),
                                          httpx.Response(500, text="down", request=httpx.Request("GET", url)))[1]
    r = await call("ffiec_call_report", {"fi_id": "480228", "period": period, "items": ["RIAD4340"]})
    n = len(calls)
    await call("ffiec_call_report", {"fi_id": "480228", "period": period, "items": ["RIAD4340"]})
    server.httpx.get = real_get
    check("failed check still serves cache, then backs off", r["items"][0]["value"] == "3500" and len(calls) == n, r)

    # files cached by v0.2.x (no catalog row) are adopted and covered by the next check
    legacy = server._cache_path("03/31/2026", "480228", "ID_RSSD", "PDF")
    legacy.write_bytes(b"%PDF-old")
    os.utime(legacy, (datetime(2026, 2, 1).timestamp(),) * 2)
    SUBMISSIONS.append({"ID_RSSD": 480228, "DateTime": "3/1/2026 9:00:00 AM"})
    db._w("DELETE FROM checks")
    await call("ffiec_download_call_report", {"fi_id": "480228", "period": "2026Q1"})
    check("legacy file amended since its mtime is refreshed", legacy.read_bytes() == b"%PDF-1.4", legacy.read_bytes()[:10])
    SUBMISSIONS.clear()

    # ---- lookups survive a restart
    server._panel_cache.clear()
    n = len(calls)
    r = await call("ffiec_find_bank", {"query": "test bank"})
    await call("ffiec_reporting_periods", {})
    check("panel and periods persisted across restart", r["matches"] == 1 and len(calls) == n, [c[0] for c in calls[n:]])
    r = await call("ffiec_reporting_periods", {"refresh": True})
    check("refresh bypasses persisted periods", calls[-1][0] == "RetrieveReportingPeriods", calls[-1][0])

    # ---- bulk import
    bulk = TMP / "FFIEC CDR Call Bulk All Schedules 09302025.zip"
    por = ("IDRSSD\tFDIC Certificate Number\tOCC Charter Number\tOTS Docket Number\tPrimary ABA Routing Number\t"
           "Financial Institution Name\tFinancial Institution City\tFinancial Institution State\t"
           "Financial Institution Filing Type\tLast Date/Time Submission Updated On\n"
           "777\t4444\t0\t0\t0\tBULK BANK\tAUSTIN\tTX\t051\t2025-10-20T10:00:00\n"
           "888\t5555\t0\t0\t0\tOTHER BANK\tDALLAS\tTX\t041\t2025-10-21T10:00:00\n")
    rc = ('"IDRSSD"\t"RCON2170"\t"RCON2200"\t"RCON3210"\n'
          '""\t"TOTAL ASSETS"\t"TOTAL DEPOSITS"\t"TOTAL EQUITY"\n'
          '777\t500000\t400000\t50000\n888\t900000\t\t90000\n')
    ri = ("IDRSSD\tRIAD4340\tRIAD4074\tRCON2170\n\tNET INCOME\tNII\tdup\n777\t4200\t15000\t500000\n888\t7000\t\t900000\n")
    with zipfile.ZipFile(bulk, "w") as z:
        z.writestr("FFIEC CDR Call Bulk POR 09302025.txt", por)
        z.writestr("FFIEC CDR Call Schedule RC 09302025.txt", rc)
        z.writestr("FFIEC CDR Call Schedule RI 09302025(1 of 2).txt", ri)
        z.writestr("Readme.txt", "notes")
    server.TIME_BUDGET = 0
    statuses = []
    for _ in range(6):
        r = await call("ffiec_import_bulk", {"path": str(bulk)})
        statuses.append(r.get("status"))
        if r.get("status") != "in_progress":
            break
    server.TIME_BUDGET = 40
    check("bulk import resumes then completes", statuses[0] == "in_progress" and statuses[-1] == "complete"
          and r["periods"] == {"09/30/2025": 2}, (statuses, r))
    n = len(calls)
    r = await call("ffiec_call_report", {"fi_id": "777", "period": "2025Q3"})
    vals = {i["mdrm"]: (i["value"], i["schedule"]) for i in r["items"]}
    check("bulk bank served locally (only the first amendment check)", r["source"] == "bulk import"
          and [c[0] for c in calls[n:]] == ["RetrieveFilersSubmissionDateTime"], [c[0] for c in calls[n:]])
    n = len(calls)
    check("bulk values, schedules, duplicate kept once", vals.get("RIAD4340") == ("4200", "RI")
          and vals.get("RCON2170") == ("500000", "RC") and len(vals) == 5, vals)
    r = await call("ffiec_bank_summary", {"fi_ids": ["888"], "period": "2025Q3"})
    check("summary from bulk (blank cells skipped)", r["banks"]["888"]["total_assets"] == 900000
          and r["banks"]["888"]["total_deposits"] is None, r)
    r = await call("ffiec_call_report", {"fi_id": "4444", "period": "2025Q3", "id_type": "FDICCertNumber", "items": ["2170"]})
    check("bulk found by FDIC cert via imported panel", r.get("source") == "bulk import" and len(calls) == n, r)
    r = await call("ffiec_cache_status", {})
    check("cache status lists bulk import", r["periods"]["09/30/2025"]["bulk_banks"] == 2 and r["bulk_imports"], r)
    check("staging database removed", not (TMP / ".ffiec_cache" / "bulk_stage.sqlite").exists())

    # an amendment after the bulk snapshot sends that bank back to the API
    YTD["09/30/2025"] = (3100, 15000)
    SUBMISSIONS.append({"ID_RSSD": 777, "DateTime": "11/2/2025 1:00:00 PM"})
    db._w("DELETE FROM checks WHERE period='09/30/2025'")
    n = len(calls)
    r = await call("ffiec_call_report", {"fi_id": "777", "period": "2025Q3", "items": ["RIAD4340"]})
    check("amended bulk bank re-fetched from API", r["source"] != "bulk import" and r["items"][0]["value"] == "3100", r)
    r = await call("ffiec_call_report", {"fi_id": "888", "period": "2025Q3", "items": ["RIAD4340"]})
    check("unamended bulk bank still local", r["source"] == "bulk import", r)
    SUBMISSIONS.clear()

    h = calls[0][1]
    check("auth headers", h["Authentication"] == "Bearer test-token" and h["UserID"] == "tester", h)
    print(f"\n{len(failures)} failure(s)")
    sys.exit(1 if failures else 0)


asyncio.run(main())
